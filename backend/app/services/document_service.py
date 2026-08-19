import os
from typing import Dict, Any
from app.services.llm_router import llm_router

class DocumentService:
    async def process_document(self, file_path: str, filename: str) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        extracted_text = ""

        try:
            if ext == ".pdf":
                extracted_text = self._parse_pdf(file_path)
            elif ext in [".docx", ".doc"]:
                extracted_text = self._parse_docx(file_path)
            elif ext in [".csv", ".xlsx", ".xls"]:
                extracted_text = self._parse_excel(file_path)
            elif ext in [".pptx", ".ppt"]:
                extracted_text = self._parse_pptx(file_path)
            else:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
        except Exception as e:
            extracted_text = f"Document content for {filename}. (Parser notice: {str(e)})"

        summary = await self._generate_ai_summary(filename, extracted_text)
        return {
            "filename": filename,
            "char_count": len(extracted_text),
            "text_preview": extracted_text[:1200],
            "summary": summary
        }

    def _parse_pdf(self, file_path: str) -> str:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
        return text if text else "PDF text extracted."

    def _parse_docx(self, file_path: str) -> str:
        import docx
        doc = docx.Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text])

    def _parse_excel(self, file_path: str) -> str:
        import pandas as pd
        df = pd.read_csv(file_path) if file_path.endswith(".csv") else pd.read_excel(file_path)
        return f"Table Dimensions: {df.shape[0]} Rows x {df.shape[1]} Columns\nHeaders: {list(df.columns)}\n\nSample Rows:\n" + df.head(10).to_string()

    def _parse_pptx(self, file_path: str) -> str:
        from pptx import Presentation
        prs = Presentation(file_path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)

    async def _generate_ai_summary(self, filename: str, text: str) -> str:
        prompt = f"Analyze and provide a comprehensive AI summary of the document '{filename}'. Include key takeaways, main topics, and document structure:\n\n{text[:3000]}"
        messages = [{"role": "user", "content": prompt}]
        return await llm_router.generate_response(messages, model="gemini-flash")

document_service = DocumentService()
